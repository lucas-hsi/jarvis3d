"""
Documentos do usuário: LER (PDF, Word, Excel, texto), CRIAR (Word, Excel, Markdown/texto) e BUSCAR por conteúdo nas pastas liberadas (índice SQLite FTS5 local, incremental).
Tudo fica nas pastas liberadas do config; arquivos novos nunca sobrescrevem nada (nome livre).
"""

import csv
import io
import json
import os
import re
import sqlite3
import time
from pathlib import Path

import arquivos
import caminhos

EXT_TEXTO = {".txt", ".md", ".csv", ".json", ".log", ".py", ".js", ".html", ".css", ".yml", ".yaml", ".ini", ".xml", ".tsv"}
EXT_INDEXAR = EXT_TEXTO | {".pdf", ".docx", ".xlsx"}
IGNORAR_DIRS = {"node_modules", ".git", "__pycache__", ".venv", "venv", "AppData", "$RECYCLE.BIN", "System Volume Information"}
TAM_MAX = 8 * 1024 * 1024
DB = caminhos.DADOS / "memoria" / "docs.db"

DEFINICOES = [
    {"name": "ler_documento",
     "description": "Lê o TEXTO de um documento do usuário em pasta liberada: PDF, Word (.docx), Excel (.xlsx), CSV ou texto. Para PDF grande use 'pagina_inicial'. Devolve o conteúdo (cortado) para você resumir, comparar ou responder.",
     "input_schema": {"type": "object", "properties": {"arquivo": {"type": "string", "description": "caminho completo ou nome do arquivo dentro de Downloads, Documentos ou Área de Trabalho"},
                                                        "pagina_inicial": {"type": "integer", "description": "PDF: página para começar (1 = primeira)"}}, "required": ["arquivo"]}},
    {"name": "criar_documento",
     "description": ("Cria um arquivo NOVO em pasta liberada: 'docx' (Word; no conteúdo use '# Título', '## Subtítulo', '- item' e parágrafos), 'xlsx' (Excel; conteúdo em CSV: uma linha por linha, vírgula ou ponto e vírgula) "
                     "ou 'md'/'txt'. Nunca sobrescreve: se o nome existir, cria outro. Depois diga onde salvou."),
     "input_schema": {"type": "object", "properties": {"tipo": {"type": "string", "enum": ["docx", "xlsx", "md", "txt"]}, "nome": {"type": "string", "description": "nome sem extensão"},
                                                        "conteudo": {"type": "string"}, "pasta": {"type": "string", "description": "pasta liberada (padrão: Documentos)"}}, "required": ["tipo", "nome", "conteudo"]}},
    {"name": "buscar_documentos",
     "description": "Procura PELO CONTEÚDO nos documentos das pastas liberadas (PDF, Word, Excel, textos) e devolve os trechos com o caminho do arquivo. Use para 'onde está o contrato com...', 'acha aquele arquivo que fala de...'.",
     "input_schema": {"type": "object", "properties": {"consulta": {"type": "string"}}, "required": ["consulta"]}},
]
NOMES = {d["name"] for d in DEFINICOES}


# ---------------------------------------------------------------- leitura
def extrair_texto(caminho: Path, pagina_inicial: int = 1, limite: int = 200_000) -> str:
    ext = caminho.suffix.lower()
    if ext == ".pdf":
        from pypdf import PdfReader
        leitor = PdfReader(str(caminho))
        partes = []
        for i, pg in enumerate(leitor.pages[max(0, pagina_inicial - 1):], start=max(1, pagina_inicial)):
            partes.append(f"[página {i} de {len(leitor.pages)}]\n" + (pg.extract_text() or ""))
            if sum(len(p) for p in partes) > limite:
                break
        return "\n".join(partes)
    if ext == ".docx":
        from docx import Document
        d = Document(str(caminho))
        linhas = [p.text for p in d.paragraphs if p.text.strip()]
        for t in d.tables:
            for r in t.rows:
                linhas.append(" | ".join(c.text.strip() for c in r.cells))
        return "\n".join(linhas)[:limite]
    if ext == ".xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(str(caminho), read_only=True, data_only=True)
        saida = []
        for ws in wb.worksheets[:6]:
            saida.append(f"[planilha '{ws.title}']")
            for n, linha in enumerate(ws.iter_rows(values_only=True)):
                if n >= 300:
                    saida.append("[... mais linhas ...]")
                    break
                saida.append("\t".join("" if c is None else str(c) for c in linha[:40]))
        return "\n".join(saida)[:limite]
    if ext in EXT_TEXTO:
        return caminho.read_text(encoding="utf-8", errors="replace")[:limite]
    raise ValueError(f"Não sei ler arquivos '{ext}'. Tipos: PDF, Word, Excel, CSV e texto.")


def ler(entrada: dict, config: dict) -> str:
    rs = arquivos.raizes(config)
    caminho = arquivos.resolver(str(entrada.get("arquivo") or ""), rs, permitir_raiz=False)
    if not caminho.is_file():
        raise ValueError("Esse arquivo não existe.")
    if caminho.stat().st_size > 40 * 1024 * 1024:
        raise ValueError("Arquivo grande demais (máx. 40 MB).")
    try:
        pg = int(entrada.get("pagina_inicial") or 1)
    except (TypeError, ValueError):
        pg = 1
    texto = extrair_texto(caminho, pg, 60_000)
    corte = " [texto cortado: peça a continuação por página]" if len(texto) >= 59_000 else ""
    return f"Arquivo: {caminho.name} ({caminho.stat().st_size // 1024} KB){corte}\n\n{texto[:30_000]}"


# ---------------------------------------------------------------- criação
def _nome_livre(pasta: Path, base: str, ext: str) -> Path:
    base = re.sub(r'[\\/:*?"<>|\r\n\t]+', " ", base).strip()[:80] or "documento"
    alvo = pasta / f"{base}.{ext}"
    n = 2
    while alvo.exists():
        alvo = pasta / f"{base} ({n}).{ext}"
        n += 1
    return alvo


def criar(entrada: dict, config: dict) -> str:
    rs = arquivos.raizes(config)
    pasta = arquivos.resolver(str(entrada.get("pasta") or "documentos"), rs) if rs else None
    if pasta is None or not pasta.is_dir():
        raise ValueError("Pasta inválida ou fora das liberadas.")
    tipo = str(entrada.get("tipo") or "").lower()
    conteudo = str(entrada.get("conteudo") or "")
    if tipo not in ("docx", "xlsx", "md", "txt") or not conteudo.strip() or len(conteudo) > 300_000:
        raise ValueError("Tipo ou conteúdo inválido.")
    alvo = _nome_livre(pasta, str(entrada.get("nome") or "documento"), tipo)
    if tipo == "docx":
        from docx import Document
        d = Document()
        for linha in conteudo.splitlines():
            t = linha.rstrip()
            if t.startswith("### "): d.add_heading(t[4:], 3)
            elif t.startswith("## "): d.add_heading(t[3:], 2)
            elif t.startswith("# "): d.add_heading(t[2:], 1)
            elif re.match(r"^[-*•]\s+", t): d.add_paragraph(re.sub(r"^[-*•]\s+", "", t), style="List Bullet")
            elif re.match(r"^\d+[.)]\s+", t): d.add_paragraph(re.sub(r"^\d+[.)]\s+", "", t), style="List Number")
            elif t: d.add_paragraph(t)
        d.save(str(alvo))
    elif tipo == "xlsx":
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        amostra = conteudo[:2000]
        delim = ";" if amostra.count(";") > amostra.count(",") else ("\t" if amostra.count("\t") > amostra.count(",") else ",")
        for linha in csv.reader(io.StringIO(conteudo), delimiter=delim):
            ws.append([_numero(c) for c in linha])
        wb.save(str(alvo))
    else:
        alvo.write_text(conteudo, encoding="utf-8")
    return f"Criei '{alvo.name}' em {alvo.parent}."


def _numero(c: str):
    t = c.strip()
    try:
        return int(t) if re.fullmatch(r"-?\d+", t) else float(t.replace(",", ".")) if re.fullmatch(r"-?\d+[.,]\d+", t) else c
    except ValueError:
        return c


# ---------------------------------------------------------------- índice de busca (FTS5, incremental)
def _db() -> sqlite3.Connection:
    DB.parent.mkdir(parents=True, exist_ok=True)
    c = sqlite3.connect(str(DB), timeout=20)
    c.execute("CREATE TABLE IF NOT EXISTS arq(caminho TEXT PRIMARY KEY, mtime REAL, tam INTEGER)")
    c.execute("CREATE VIRTUAL TABLE IF NOT EXISTS trechos USING fts5(caminho UNINDEXED, texto, tokenize='unicode61 remove_diacritics 2')")
    return c


def indexar(config: dict, orcamento_s: float = 25.0) -> dict:
    """Indexa (só o que mudou) os documentos das pastas liberadas. Pára ao estourar o tempo; continua na próxima."""
    t0 = time.monotonic()
    c = _db()
    vistos, novos, pulados = set(), 0, 0
    try:
        for raiz in arquivos.raizes(config):
            for pasta, dirs, nomes in os.walk(raiz):
                dirs[:] = [d for d in dirs if d not in IGNORAR_DIRS and not d.startswith(".")]
                for nome in nomes:
                    p = Path(pasta) / nome
                    if p.suffix.lower() not in EXT_INDEXAR:
                        continue
                    try:
                        st = p.stat()
                    except OSError:
                        continue
                    if st.st_size > TAM_MAX or len(vistos) >= 6000:
                        pulados += 1
                        continue
                    chave = str(p)
                    vistos.add(chave)
                    ja = c.execute("SELECT mtime FROM arq WHERE caminho=?", (chave,)).fetchone()
                    if ja and abs(ja[0] - st.st_mtime) < 1:
                        continue
                    if time.monotonic() - t0 > orcamento_s:
                        return {"novos": novos, "arquivos": len(vistos), "parcial": True}
                    try:
                        texto = extrair_texto(p, 1, 150_000)
                    except Exception:
                        texto = ""
                    c.execute("DELETE FROM trechos WHERE caminho=?", (chave,))
                    for i in range(0, len(texto), 1200):
                        pedaco = texto[i:i + 1400].strip()
                        if pedaco:
                            c.execute("INSERT INTO trechos(caminho, texto) VALUES (?, ?)", (chave, pedaco))
                    c.execute("INSERT OR REPLACE INTO arq VALUES (?, ?, ?)", (chave, st.st_mtime, st.st_size))
                    novos += 1
        for (chave,) in c.execute("SELECT caminho FROM arq").fetchall():       # arquivos que sumiram saem do índice
            if chave not in vistos and not Path(chave).exists():
                c.execute("DELETE FROM arq WHERE caminho=?", (chave,))
                c.execute("DELETE FROM trechos WHERE caminho=?", (chave,))
        c.commit()
        return {"novos": novos, "arquivos": len(vistos), "parcial": False, "pulados": pulados}
    finally:
        c.commit()
        c.close()


def buscar(consulta: str, config: dict, n: int = 8) -> str:
    termos = [t for t in re.findall(r"\w{3,}", consulta.lower())][:8]
    if not termos:
        raise ValueError("Diga o que procurar.")
    indexar(config, 6.0)                                   # pega o que mudou desde a última vez (rápido; o resto já estava indexado)
    c = _db()
    try:
        q = " OR ".join(f'"{t}"' for t in termos)
        linhas = c.execute("SELECT caminho, snippet(trechos, 1, '«', '»', ' … ', 28), bm25(trechos) FROM trechos WHERE trechos MATCH ? ORDER BY bm25(trechos) LIMIT 40", (q,)).fetchall()
    finally:
        c.close()
    if not linhas:
        return f"Não achei nada sobre '{consulta}' nos documentos das pastas liberadas."
    vistos, saida = set(), []
    for caminho, trecho, _ in linhas:
        if caminho in vistos:
            continue
        vistos.add(caminho)
        saida.append(f"- {caminho}\n  {trecho}")
        if len(saida) >= n:
            break
    return f"Documentos que falam de '{consulta}':\n" + "\n".join(saida)
